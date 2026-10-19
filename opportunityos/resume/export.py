import docx
from docx.shared import Pt

# Writes the tailored draft to a NEW .docx; the master resume file is never touched.
def export_docx(resume, path):
    d = docx.Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for i, line in enumerate(resume["header"]):
        p = d.add_paragraph(); r = p.add_run(line); r.bold = i == 0
        if i == 0:
            r.font.size = Pt(16)
    def heading(t):
        p = d.add_paragraph(); r = p.add_run(t.upper()); r.bold = True; r.font.size = Pt(11)
    if resume["education"]:
        heading("Education")
        for l in resume["education"]:
            d.add_paragraph(l)
    if resume["skills"]:
        heading("Skills"); d.add_paragraph(", ".join(resume["skills"]))
    for sec, title in (("experience", "Experience"), ("projects", "Projects")):
        if resume[sec]:
            heading(title)
            for item in resume[sec]:
                p = d.add_paragraph(); p.add_run(item["name"]).bold = True
                for b in item["bullets"]:
                    d.add_paragraph(b, style="List Bullet")
    if resume["certifications"]:
        heading("Certifications")
        for l in resume["certifications"]:
            d.add_paragraph(l)
    d.add_paragraph("DRAFT - review every line before submitting.").runs[0].italic = True
    d.save(path)

def gap_report_md(jd, report_rows, changes, decisions, issues):
    L = [f"# Gap report: {jd.get('title') or 'job description'}", "", "## Requirements", "",
         "| Requirement | Type | Resume evidence | Final status |", "|---|---|---|---|"]
    for r in report_rows:
        L.append(f"| {r['requirement']} | {r['kind']} | {r['status']} | {r['final']} |")
    L += ["", "## Change review", ""]
    for c in changes:
        L.append(f"- **{c['type']}** ({decisions.get(c['id'], 'pending')}): {c['reason']}. Source: {c['source']}.")
        if c["original"]:
            L.append(f"  - Before: {c['original']}")
        L.append(f"  - Proposed: {c['proposed']}")
    L += ["", "## Truthfulness checks", ""]
    L += [f"- {i['severity'].upper()}: {i['check']} - {i['detail']}" for i in issues] or ["- No issues found."]
    return "\n".join(L)
